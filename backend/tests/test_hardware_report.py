import unittest
from io import BytesIO
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.main import app
from app.services.hardware_report_service import (
    build_hardware_pdf,
    build_hardware_xlsx,
)


SAMPLE_ANALYSIS = {
    "title": "Báo cáo phân tích phần cứng",
    "summary": "Tóm tắt chất lượng sản xuất.",
    "data_scope": {"has_model_data": True},
    "kpis": [
        {
            "name": "Tổng sản lượng",
            "value": 1200,
            "unit": "sản phẩm",
            "description": "Tổng sản lượng đã ghi nhận.",
        },
        {
            "name": "Tỷ lệ lỗi chung",
            "value": 2.5,
            "unit": "%",
            "description": "Tỷ lệ lỗi.",
        },
    ],
    "daily_production": [
        {
            "date": "2026-10-01",
            "production": 1200,
            "defects": 30,
            "defect_rate": 2.5,
            "downtime": 5,
        }
    ],
    "product_performance": [{"product": "Widget", "production": 1200}],
    "model_performance": [{"model": "Model A", "production": 1200}],
    "component_performance": [{"component": "Board", "production": 1200}],
    "supplier_performance": [{"supplier": "Supplier X", "defects": 30}],
    "batch_performance": [{"batch": "Batch 1", "defects": 30}],
    "defect_type_analysis": [{"defect_type": "Scratch", "defects": 30}],
    "data_quality": {"quality_status": "good", "quality_score": 95},
    "data_processing": {"rows_used_for_kpi": 1},
    "insights": ["Theo dõi tỷ lệ lỗi."],
    "warnings": ["Một số trường còn thiếu."],
    "recommendations": ["Kiểm tra lại dữ liệu đầu vào."],
    "ai_analysis": {
        "executive_summary": {
            "summary": "Sản xuất ổn định.",
            "key_observations": ["Có dữ liệu đủ dùng."],
        },
        "risk_areas": [],
        "recommended_actions": [],
    },
}


class HardwareReportServiceTests(unittest.TestCase):
    def test_xlsx_contains_ten_report_sheets_and_analysis_data(self):
        workbook_bytes = build_hardware_xlsx(SAMPLE_ANALYSIS)
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Executive Summary",
                "KPI Overview",
                "Production Analysis",
                "Quality Analysis",
                "Supplier Analysis",
                "Batch Analysis",
                "Defect Analysis",
                "AI Insights",
                "Warnings",
                "Recommendations",
            ],
        )
        kpi_values = [
            cell.value
            for row in workbook["KPI Overview"].iter_rows()
            for cell in row
        ]
        self.assertIn("Tổng sản lượng", kpi_values)
        self.assertIn(1200, kpi_values)
        self.assertIn("No data provided by analysis", [
            cell.value
            for row in workbook["AI Insights"].iter_rows()
            for cell in row
        ])

    def test_xlsx_escapes_formula_like_user_values(self):
        workbook = load_workbook(
            BytesIO(
                build_hardware_xlsx(
                    {
                        "title": "Report",
                        "insights": ["=HYPERLINK(\"https://example.com\")"],
                    }
                )
            ),
            data_only=True,
        )
        insight_values = [
            cell.value
            for row in workbook["AI Insights"].iter_rows()
            for cell in row
        ]

        self.assertIn(
            "'=HYPERLINK(\"https://example.com\")",
            insight_values,
        )

    def test_pdf_has_pdf_signature_and_vietnamese_report_text(self):
        pdf_bytes = build_hardware_pdf(SAMPLE_ANALYSIS)
        extracted_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(BytesIO(pdf_bytes)).pages
        )

        self.assertTrue(pdf_bytes.startswith(b"%PDF-"))
        self.assertIn("HardwareAI Report", extracted_text)
        self.assertIn("Tóm tắt chất lượng sản xuất.", extracted_text)
        self.assertIn("Recommendations", extracted_text)


class HardwareReportEndpointTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_export_endpoint_serves_pdf_and_xlsx(self):
        pdf_response = self.client.post(
            "/api/hardware/export?format=pdf",
            json=SAMPLE_ANALYSIS,
        )
        xlsx_response = self.client.post(
            "/api/hardware/export?format=xlsx",
            json=SAMPLE_ANALYSIS,
        )

        self.assertEqual(pdf_response.status_code, 200)
        self.assertEqual(pdf_response.headers["content-type"], "application/pdf")
        self.assertIn(
            'filename="HardwareAI_Report.pdf"',
            pdf_response.headers["content-disposition"],
        )
        self.assertTrue(pdf_response.content.startswith(b"%PDF-"))

        self.assertEqual(xlsx_response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_response.headers["content-type"],
        )
        self.assertIn(
            'filename="HardwareAI_Report.xlsx"',
            xlsx_response.headers["content-disposition"],
        )
        with ZipFile(BytesIO(xlsx_response.content)) as xlsx_file:
            self.assertIn("xl/workbook.xml", xlsx_file.namelist())

    def test_export_endpoint_rejects_empty_or_unrelated_payloads(self):
        for payload in ({}, {"unrelated": "data"}):
            with self.subTest(payload=payload):
                response = self.client.post(
                    "/api/hardware/export",
                    json=payload,
                )
                self.assertEqual(response.status_code, 422)

    def test_export_endpoint_rejects_unsupported_format(self):
        response = self.client.post(
            "/api/hardware/export?format=doc",
            json=SAMPLE_ANALYSIS,
        )

        self.assertEqual(response.status_code, 422)


if __name__ == "__main__":
    unittest.main()
