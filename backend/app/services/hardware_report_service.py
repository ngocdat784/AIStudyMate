from __future__ import annotations

import json
import math
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as PDFTable,
    TableStyle,
)


REPORT_SECTIONS = (
    ("Executive Summary", "executive_summary"),
    ("KPI Overview", "kpi_overview"),
    ("Production Analysis", "production_analysis"),
    ("Quality Analysis", "quality_analysis"),
    ("Supplier Analysis", "supplier_analysis"),
    ("Batch Analysis", "batch_analysis"),
    ("Defect Analysis", "defect_analysis"),
    ("AI Insights", "ai_insights"),
    ("Warnings", "warnings"),
    ("Recommendations", "recommendations"),
)
_PDF_FONT_LOCK = Lock()


def _safe_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, (str, int, float)):
        return value
    return str(value)


def _display_value(value: Any) -> str:
    safe = _safe_value(value)
    if isinstance(safe, float):
        return f"{safe:,.2f}".rstrip("0").rstrip(".")
    return str(safe)


def _flatten_mapping(
    value: Any,
    prefix: str = "",
) -> list[tuple[str, Any]]:
    if isinstance(value, dict):
        rows: list[tuple[str, Any]] = []
        for key, nested_value in value.items():
            nested_prefix = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten_mapping(nested_value, nested_prefix))
        return rows
    if isinstance(value, list):
        if not value:
            return [(prefix, "No data provided by analysis")]
        if all(not isinstance(item, (dict, list)) for item in value):
            return [(prefix, "; ".join(_display_value(item) for item in value))]
        rows = []
        for index, item in enumerate(value, start=1):
            rows.extend(_flatten_mapping(item, f"{prefix}[{index}]"))
        return rows
    return [(prefix, value)]


def _records(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [
            item if isinstance(item, dict) else {"Item": item}
            for item in value
        ]
    if isinstance(value, dict):
        return [{"Field": key, "Value": item} for key, item in value.items()]
    return []


def _section_blocks(
    payload: dict[str, Any],
    section_key: str,
) -> list[tuple[str, list[dict[str, Any]]]]:
    if section_key == "executive_summary":
        return [
            (
                "Summary",
                [
                    {"Field": "Title", "Value": payload.get("title")},
                    {"Field": "Summary", "Value": payload.get("summary")},
                ],
            ),
            (
                "Data Scope",
                [
                    {"Field": key, "Value": value}
                    for key, value in _flatten_mapping(
                        payload.get("data_scope", {})
                    )
                ],
            ),
        ]
    if section_key == "kpi_overview":
        return [("KPIs", _records(payload.get("kpis")))]
    if section_key == "production_analysis":
        return [
            ("Daily Production", _records(payload.get("daily_production"))),
            ("Product Performance", _records(payload.get("product_performance"))),
            ("Model Performance", _records(payload.get("model_performance"))),
            (
                "Component Performance",
                _records(payload.get("component_performance")),
            ),
        ]
    if section_key == "quality_analysis":
        quality_kpis = [
            item
            for item in _records(payload.get("kpis"))
            if any(
                term in str(item.get("name", "")).casefold()
                for term in ("quality", "defect", "lỗi", "chất lượng")
            )
        ]
        return [
            ("Quality KPIs", quality_kpis),
            (
                "Data Quality",
                [
                    {"Field": key, "Value": value}
                    for key, value in _flatten_mapping(
                        payload.get("data_quality", {})
                    )
                ],
            ),
            (
                "Data Processing",
                [
                    {"Field": key, "Value": value}
                    for key, value in _flatten_mapping(
                        payload.get("data_processing", {})
                    )
                ],
            ),
        ]
    if section_key == "supplier_analysis":
        return [("Supplier Performance", _records(payload.get("supplier_performance")))]
    if section_key == "batch_analysis":
        return [("Batch Performance", _records(payload.get("batch_performance")))]
    if section_key == "defect_analysis":
        return [
            ("Defect Types", _records(payload.get("defect_type_analysis"))),
            (
                "AI Risk Areas",
                [
                    item
                    for item in _records(
                        (
                            payload.get("ai_analysis", {}).get("risk_areas")
                            if isinstance(payload.get("ai_analysis"), dict)
                            else None
                        )
                    )
                    if item.get("entity_type") == "defect_type"
                ],
            ),
        ]
    if section_key == "ai_insights":
        insight_items = payload.get("insights")
        rows = (
            [{"Insight": item} for item in insight_items]
            if isinstance(insight_items, list)
            else []
        )
        rows.extend(
            {"Field": key, "Value": value}
            for key, value in _flatten_mapping(payload.get("ai_analysis", {}))
        )
        return [("Insights and AI Analysis", rows)]
    if section_key == "warnings":
        return [("Warnings", _records(payload.get("warnings")))]
    if section_key == "recommendations":
        recommendation_items = payload.get("recommendations")
        rows = (
            [{"Recommendation": item} for item in recommendation_items]
            if isinstance(recommendation_items, list)
            else []
        )
        ai_analysis = payload.get("ai_analysis")
        if isinstance(ai_analysis, dict):
            rows.extend(
                _records(ai_analysis.get("recommended_actions"))
            )
        return [("Recommended Actions", rows)]
    return []


def _table_rows(records: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    if not records or not any(records):
        return ["Information"], [["No data provided by analysis"]]
    headers = list(dict.fromkeys(
        key for record in records for key in record
    ))
    rows = [
        [_safe_value(record.get(header)) for header in headers]
        for record in records
    ]
    return headers, rows


def _xlsx_cell_value(value: Any) -> Any:
    safe = _safe_value(value)
    if isinstance(safe, str) and safe.startswith(("=", "+", "-", "@")):
        return f"'{safe}"
    return safe


def build_hardware_xlsx(payload: dict[str, Any]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for section_title, section_key in REPORT_SECTIONS:
        worksheet = workbook.create_sheet(section_title)
        worksheet.sheet_view.showGridLines = False
        worksheet["A1"] = section_title
        worksheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        worksheet["A1"].fill = PatternFill("solid", fgColor="17365D")
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

        current_row = 3
        table_number = 0
        for block_title, records in _section_blocks(payload, section_key):
            worksheet.cell(row=current_row, column=1, value=block_title)
            worksheet.cell(row=current_row, column=1).font = Font(
                bold=True,
                size=12,
                color="17365D",
            )
            current_row += 1
            headers, rows = _table_rows(records)
            header_row = current_row
            for column, header in enumerate(headers, start=1):
                cell = worksheet.cell(
                    row=current_row,
                    column=column,
                    value=_xlsx_cell_value(header),
                )
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4472C4")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            current_row += 1
            for row in rows:
                for column, value in enumerate(row, start=1):
                    cell = worksheet.cell(
                        row=current_row,
                        column=column,
                        value=_xlsx_cell_value(value),
                    )
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                current_row += 1

            table_number += 1
            if rows:
                table = Table(
                    displayName=f"ReportTable{len(workbook.worksheets)}_{table_number}",
                    ref=(
                        f"A{header_row}:"
                        f"{worksheet.cell(row=header_row + len(rows), column=len(headers)).coordinate}"
                    ),
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                worksheet.add_table(table)
            current_row += 2

        for column_index in range(1, worksheet.max_column + 1):
            column_letter = get_column_letter(column_index)
            column_cells = worksheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
            )
            column_values = next(column_cells)
            max_width = max(
                (
                    min(len(str(cell.value or "")), 60)
                    for cell in column_values
                ),
                default=12,
            )
            worksheet.column_dimensions[column_letter].width = max(
                12,
                min(max_width + 2, 62),
            )
        worksheet.freeze_panes = "A4"
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_font_names() -> tuple[str, str]:
    regular_name = "HardwareAIArial"
    bold_name = "HardwareAIArialBold"
    with _PDF_FONT_LOCK:
        registered = set(pdfmetrics.getRegisteredFontNames())
        if regular_name in registered and bold_name in registered:
            return regular_name, bold_name

        font_pairs = (
            (
                Path(r"C:\Windows\Fonts\arial.ttf"),
                Path(r"C:\Windows\Fonts\arialbd.ttf"),
            ),
            (
                Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            ),
        )
        for regular_path, bold_path in font_pairs:
            if regular_path.is_file() and bold_path.is_file():
                if regular_name not in registered:
                    pdfmetrics.registerFont(
                        TTFont(regular_name, str(regular_path))
                    )
                if bold_name not in registered:
                    pdfmetrics.registerFont(
                        TTFont(bold_name, str(bold_path))
                    )
                return regular_name, bold_name
    raise RuntimeError(
        "Không tìm thấy font Arial hoặc DejaVu Sans để tạo PDF tiếng Việt."
    )


def _pdf_paragraph(value: Any, style: ParagraphStyle) -> Paragraph:
    text = escape(_display_value(value)).replace("\n", "<br/>")
    return Paragraph(text or "—", style)


def build_hardware_pdf(payload: dict[str, Any]) -> bytes:
    regular_font, bold_font = _pdf_font_names()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="HardwareAI Report",
        author="HardwareAI",
    )
    base_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "HardwareReportTitle",
        parent=base_styles["Title"],
        fontName=bold_font,
        fontSize=22,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#17365D"),
        spaceAfter=6 * mm,
    )
    section_style = ParagraphStyle(
        "HardwareReportSection",
        parent=base_styles["Heading2"],
        fontName=bold_font,
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#17365D"),
        spaceBefore=5 * mm,
        spaceAfter=2 * mm,
    )
    subsection_style = ParagraphStyle(
        "HardwareReportSubsection",
        parent=base_styles["Heading3"],
        fontName=bold_font,
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#4472C4"),
        spaceBefore=2 * mm,
        spaceAfter=1 * mm,
    )
    cell_style = ParagraphStyle(
        "HardwareReportCell",
        parent=base_styles["BodyText"],
        fontName=regular_font,
        fontSize=7,
        leading=9,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        "HardwareReportHeader",
        parent=cell_style,
        fontName=bold_font,
        textColor=colors.white,
    )

    story = [
        Paragraph("HardwareAI Report", title_style),
        Paragraph(
            escape(str(payload.get("title") or "Hardware Analysis")),
            ParagraphStyle(
                "HardwareReportSubtitle",
                parent=cell_style,
                fontName=bold_font,
                fontSize=12,
                alignment=TA_CENTER,
            ),
        ),
        Spacer(1, 4 * mm),
    ]

    for section_title, section_key in REPORT_SECTIONS:
        story.append(Paragraph(escape(section_title), section_style))
        blocks = _section_blocks(payload, section_key)
        for block_title, records in blocks:
            story.append(Paragraph(escape(block_title), subsection_style))
            headers, rows = _table_rows(records)
            table_data = [
                [_pdf_paragraph(header, header_style) for header in headers],
                *[
                    [_pdf_paragraph(value, cell_style) for value in row]
                    for row in rows
                ],
            ]
            table = PDFTable(
                table_data,
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E2F3")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                            colors.white,
                            colors.HexColor("#F3F6FA"),
                        ]),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.extend([table, Spacer(1, 2 * mm)])

    document.build(story)
    return output.getvalue()
