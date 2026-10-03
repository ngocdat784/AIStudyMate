from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader


MAX_TEXT_LENGTH = 50000
MAX_ROWS = 5000
MAX_COLUMNS = 100


# =========================================================
# TXT
# =========================================================

def extract_text_from_txt(file_bytes: bytes) -> str:
    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1258",
        "latin-1",
    ]

    for encoding in encodings:
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError("Không thể đọc file TXT.")


# =========================================================
# PDF
# =========================================================

def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(
            f"Không thể đọc file PDF: {str(e)}"
        )

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text()
        except Exception:
            text = None

        if text:
            pages.append(
                f"[Trang {page_number}]\n{text}"
            )

    return "\n\n".join(pages)


# =========================================================
# DOCX
# =========================================================

def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = Document(BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(
            f"Không thể đọc file DOCX: {str(e)}"
        )

    sections = []

    # -----------------------------------------
    # Paragraphs
    # -----------------------------------------

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # -----------------------------------------
    # Tables
    # -----------------------------------------

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):
        rows = []

        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                cells.append(cell_text)

            rows.append(" | ".join(cells))

        if rows:
            sections.append(
                f"[Bảng {table_index}]\n"
                + "\n".join(rows)
            )

    return "\n\n".join(sections)


# =========================================================
# CSV
# =========================================================

def extract_production_data_from_csv(
    file_bytes: bytes,
) -> str:

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1258",
        "latin-1",
    ]

    dataframe = None

    last_error = None

    for encoding in encodings:
        try:
            dataframe = pd.read_csv(
                BytesIO(file_bytes),
                encoding=encoding,
            )
            break

        except Exception as e:
            last_error = e

    if dataframe is None:
        raise ValueError(
            f"Không thể đọc file CSV: {str(last_error)}"
        )

    return dataframe_to_text(
        dataframe,
        source_type="CSV",
    )


# =========================================================
# XLSX
# =========================================================

def extract_production_data_from_xlsx(
    file_bytes: bytes,
) -> str:

    try:
        excel_file = pd.ExcelFile(
            BytesIO(file_bytes),
            engine="openpyxl",
        )
    except Exception as e:
        raise ValueError(
            f"Không thể đọc file Excel: {str(e)}"
        )

    sections = []

    for sheet_name in excel_file.sheet_names:

        try:
            dataframe = pd.read_excel(
                excel_file,
                sheet_name=sheet_name,
            )

        except Exception as e:
            raise ValueError(
                f"Không thể đọc sheet '{sheet_name}': {str(e)}"
            )

        if dataframe.empty:
            continue

        sheet_text = dataframe_to_text(
            dataframe,
            source_type="XLSX",
            sheet_name=sheet_name,
        )

        sections.append(sheet_text)

    if not sections:
        raise ValueError(
            "Không tìm thấy dữ liệu trong file Excel."
        )

    return "\n\n".join(sections)


# =========================================================
# DataFrame → AI-readable text
# =========================================================

def dataframe_to_text(
    dataframe: pd.DataFrame,
    source_type: str,
    sheet_name: str | None = None,
) -> str:

    dataframe = dataframe.copy()

    if dataframe.empty:
        return ""

    # -----------------------------------------
    # Giới hạn dữ liệu
    # -----------------------------------------

    dataframe = dataframe.iloc[
        :MAX_ROWS,
        :MAX_COLUMNS,
    ]

    # -----------------------------------------
    # Chuẩn hóa column names
    # -----------------------------------------

    dataframe.columns = [
        str(column).strip()
        for column in dataframe.columns
    ]

    # -----------------------------------------
    # Loại bỏ dòng hoàn toàn rỗng
    # -----------------------------------------

    dataframe = dataframe.dropna(
        how="all"
    )

    if dataframe.empty:
        return ""

    # -----------------------------------------
    # Chuyển dữ liệu thành text
    # -----------------------------------------

    lines = []

    header = f"[{source_type}]"

    if sheet_name:
        header += f" Sheet: {sheet_name}"

    lines.append(header)

    lines.append(
        f"Số dòng: {len(dataframe)}"
    )

    lines.append(
        f"Số cột: {len(dataframe.columns)}"
    )

    lines.append(
        "Các cột: "
        + ", ".join(
            str(column)
            for column in dataframe.columns
        )
    )

    lines.append("")

    # -----------------------------------------
    # Dữ liệu
    # -----------------------------------------

    for index, row in dataframe.iterrows():

        values = []

        for column in dataframe.columns:

            value = row[column]

            if pd.isna(value):
                value = ""

            values.append(
                f"{column}: {value}"
            )

        lines.append(
            f"Row {index + 1}: "
            + " | ".join(values)
        )

    return "\n".join(lines)


# =========================================================
# Main Parser
# =========================================================

def extract_text(
    file_bytes: bytes,
    filename: str,
    content_type: str | None = None,
) -> str:

    extension = Path(filename).suffix.lower()

    # =====================================================
    # TEXT DOCUMENTS
    # =====================================================

    if extension == ".txt" or content_type == "text/plain":

        text = extract_text_from_txt(
            file_bytes
        )

    elif (
        extension == ".pdf"
        or content_type == "application/pdf"
    ):

        text = extract_text_from_pdf(
            file_bytes
        )

    elif (
        extension == ".docx"
        or content_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):

        text = extract_text_from_docx(
            file_bytes
        )

    # =====================================================
    # PRODUCTION DATA
    # =====================================================

    elif extension == ".csv":

        text = extract_production_data_from_csv(
            file_bytes
        )

    elif extension == ".xlsx":

        text = extract_production_data_from_xlsx(
            file_bytes
        )

    else:

        raise ValueError(
            "Định dạng file không được hỗ trợ. "
            "TextileAI hỗ trợ PDF, DOCX, TXT, CSV và XLSX."
        )

    # =====================================================
    # Normalize
    # =====================================================

    text = text.strip()

    if not text:

        raise ValueError(
            "Không tìm thấy nội dung trong file."
        )

    # =====================================================
    # Limit text sent to AI
    # =====================================================

    if len(text) > MAX_TEXT_LENGTH:

        text = text[:MAX_TEXT_LENGTH]

        text += (
            "\n\n"
            "[Nội dung đã được giới hạn "
            "do vượt quá kích thước xử lý của TextileAI.]"
        )

    return text